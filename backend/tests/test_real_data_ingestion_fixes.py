"""Regression tests for ingestion bugs found by running the pipeline over real case data.

Each test encodes a failure that was silently discarding evidence — the pipeline reported
success while dropping rows or whole files. Fixtures reproduce the exact shapes seen in
the real FIR exports (no real data is committed).
"""

from datetime import UTC, datetime, timedelta, timezone

import pytest

from backend.app.ingestion import detector
from backend.app.ingestion import service as ing
from backend.app.ingestion.parsers import tabular

IST = timezone(timedelta(hours=5, minutes=30))


# ── Trailing-comma CSV: data rows carry one more field than the header ──────────
# pandas resolves that by promoting column 0 to the index, shifting every column left.
# In the real Vodafone-Idea CDR this put "Incoming" in the A-party phone column and a
# duration in the call-date column, so normalization rejected 42,873 of 42,873 rows.

TRAILING_COMMA_CSV = (
    "Target /A PARTY NUMBER,CALL_TYPE,B PARTY NUMBER,Call date,Call Initiation Time\n"
    "919702000558,Incoming,918141122818,01-08-2024,00:04:06,\n"
    "919702000558,Outgoing,919870386595,01-08-2024,00:04:18,\n"
)


def test_trailing_comma_csv_does_not_shift_columns(tmp_path):
    p = tmp_path / "cdr.csv"
    p.write_text(TRAILING_COMMA_CSV, encoding="utf-8")

    df = tabular.read(str(p))

    # The A-party column must hold the phone number, not the call type.
    assert df["Target /A PARTY NUMBER"].iloc[0] == "919702000558"
    assert df["CALL_TYPE"].iloc[0] == "Incoming"
    assert df["Call date"].iloc[0] == "01-08-2024"


def test_trailing_comma_csv_rows_survive_ingestion(tmp_path):
    p = tmp_path / "cdr.csv"
    p.write_text(TRAILING_COMMA_CSV, encoding="utf-8")

    pf = ing.parse_file(str(p))

    assert len(pf.records) == 2
    assert pf.records[0]["Target /A PARTY NUMBER"] == "919702000558"


# ── Duplicate / blank header names silently overwrote earlier columns ───────────

def test_duplicate_headers_are_disambiguated():
    out = ing._dedupe_headers(["Amount", "Date", "Amount", "", "Amount"])
    assert out == ["Amount", "Date", "Amount__2", "column_3", "Amount__3"]
    assert len(set(out)) == len(out)


def test_duplicate_headers_keep_every_column(tmp_path):
    p = tmp_path / "dup.csv"
    p.write_text("Amount,Amount\n10,20\n", encoding="utf-8")

    pf = ing.parse_file(str(p))

    # Both columns must survive; previously the second overwrote the first.
    assert len(pf.records) == 1
    assert sorted(v for k, v in pf.records[0].items() if k != "_provenance") == ["10", "20"]


# ── Extensions lie: sniff the container instead of trusting the suffix ──────────

def test_sniff_detects_zip_regardless_of_extension(tmp_path):
    p = tmp_path / "actually_xlsx.xls"          # xlsx content named .xls
    p.write_bytes(b"PK\x03\x04" + b"\x00" * 32)
    assert detector.sniff_container(str(p)) == "zip"
    assert detector.detect_format(str(p)) == "xlsx"


def test_sniff_detects_legacy_ole2_xls(tmp_path):
    p = tmp_path / "legacy.xls"
    p.write_bytes(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\x00" * 32)
    assert detector.sniff_container(str(p)) == "ole2"
    assert detector.detect_format(str(p)) == "xlsx"


def test_text_report_named_xls_falls_back_to_delimited(tmp_path):
    # A real Bank of Baroda export is a fixed-width text report named .xls. It must not
    # reach the Excel reader (opaque "format cannot be determined" error).
    p = tmp_path / "report.xls"
    p.write_text("\r\n\r\n 23-07-2025  BANK OF BARODA   Page 1\r\n", encoding="utf-8")
    assert detector.sniff_container(str(p)) == "text"
    assert detector.detect_format(str(p)) == "csv"


def test_docx_still_resolves_to_docx(tmp_path):
    p = tmp_path / "tables.docx"
    p.write_bytes(b"PK\x03\x04" + b"\x00" * 32)
    assert detector.detect_format(str(p)) == "docx"


def test_appledouble_is_rejected_by_format_detection(tmp_path):
    p = tmp_path / "._report.xlsx"
    p.write_bytes(b"\x00\x05\x16\x07" + b"\x00" * 32)
    assert detector.sniff_container(str(p)) == "appledouble"
    with pytest.raises(ValueError, match="AppleDouble"):
        detector.detect_format(str(p))


# ── macOS sidecars must never be walked as evidence ────────────────────────────

def test_parse_directory_skips_appledouble_and_macosx(tmp_path):
    (tmp_path / "real.csv").write_text("Amount,Date\n10,01-08-2024\n", encoding="utf-8")
    (tmp_path / "._real.csv").write_bytes(b"\x00\x05\x16\x07" + b"\x00" * 32)
    macosx = tmp_path / "__MACOSX"
    macosx.mkdir()
    (macosx / "real.csv").write_bytes(b"\x00\x05\x16\x07" + b"\x00" * 32)

    parsed = ing.parse_directory(str(tmp_path), include_pdf=False)

    names = [p.path.replace("\\", "/").rsplit("/", 1)[-1] for p in parsed]
    assert names == ["real.csv"]


# ── /v1/events sort must not mix naive and aware datetimes ─────────────────────

def test_event_sort_fallback_is_timezone_aware():
    """Event timestamps are tz-aware; a naive datetime.min fallback raises TypeError
    and turns GET /v1/events into a 500."""
    events = [
        {"timestamp_start": datetime(2026, 6, 1, 22, 44, tzinfo=IST)},
        {"timestamp_start": None},
    ]
    ordered = sorted(
        events,
        key=lambda e: e.get("timestamp_start") or datetime.min.replace(tzinfo=UTC),
    )
    assert ordered[0]["timestamp_start"] is None


# ── Profile-matching gate ──────────────────────────────────────────────────────

def test_required_any_gate_is_actually_enforced():
    """`match` lives under the `profile:` block. Reading it from the top level returned
    {} for every profile, so the gate never fired and real IPDR exports were scored as
    CDR on shared IMEI/IMSI columns alone."""
    profile = {
        "profile": {"id": "x", "source": "CDR",
                    "match": {"required_any": ["A PARTY NUMBER"]}},
        "field_map": {"imei": {"aliases": ["IMEI"]}, "imsi": {"aliases": ["IMSI"]}},
    }
    # Shares IMEI/IMSI but has no A-party -> must be rejected outright, not scored.
    assert detector.score_profile(["IMEI", "IMSI", "Public IP Address"], profile) == 0.0
    assert detector.score_profile(["A PARTY NUMBER", "IMEI"], profile) > 0.0


def test_required_all_gate_is_enforced():
    profile = {
        "profile": {"id": "x", "match": {"required_all": ["IMEI", "IMSI"]}},
        "field_map": {"imei": {"aliases": ["IMEI"]}, "imsi": {"aliases": ["IMSI"]}},
    }
    assert detector.score_profile(["IMEI"], profile) == 0.0
    assert detector.score_profile(["IMEI", "IMSI"], profile) > 0.0


def test_every_profile_gate_is_reachable_from_its_own_aliases():
    """A required_any entry that appears in no field_map alias is almost always drift:
    the gate then rejects files the profile is otherwise able to map. This caught
    cdr_vodafone_idea, whose gate omitted the Mobile_No dialect it maps."""
    from backend.app.core import config

    problems = []
    for plist in config.profiles().values():
        for prof in plist:
            gate = prof.get("profile", {}).get("match", {}).get("required_any", [])
            if not gate:
                continue
            aliases = {a.strip().lower()
                       for spec in prof.get("field_map", {}).values()
                       for a in spec.get("aliases", [])}
            unknown = [g for g in gate if g.strip().lower() not in aliases]
            if unknown:
                problems.append((prof["profile"]["id"], unknown))
    assert not problems, f"required_any entries not present in any alias list: {problems}"


# ── Archives: evidence sealed inside (nested) ZIPs ─────────────────────────────
# On the real case, 92 archives were never opened — 83 structured files plus 129 PDFs
# sat inside them, and nothing reported the omission.

def _zip_with(tmp_path, name, entries):
    import zipfile
    p = tmp_path / name
    with zipfile.ZipFile(p, "w") as zf:
        for member, body in entries.items():
            zf.writestr(member, body)
    return p


def test_archive_members_are_parsed(tmp_path):
    _zip_with(tmp_path, "bank.zip", {"stmt.csv": "Amount,Date\n10,01-08-2024\n"})
    parsed = ing.parse_directory(str(tmp_path), include_pdf=False)
    assert [p.records[0]["Amount"] for p in parsed if p.records] == ["10"]


def test_archive_provenance_names_the_container(tmp_path):
    """A statement pulled out of bank.zip must still cite the exhibit it came from."""
    _zip_with(tmp_path, "bank.zip", {"stmt.csv": "Amount,Date\n10,01-08-2024\n"})
    parsed = [p for p in ing.parse_directory(str(tmp_path), include_pdf=False) if p.records]

    assert parsed[0].container == "bank.zip"
    assert parsed[0].records[0]["_provenance"]["source_file"] == "bank.zip → stmt.csv"


def test_nested_archives_are_expanded(tmp_path):
    import io
    import zipfile
    inner = io.BytesIO()
    with zipfile.ZipFile(inner, "w") as zf:
        zf.writestr("deep.csv", "Amount,Date\n99,01-08-2024\n")
    _zip_with(tmp_path, "outer.zip", {"inner.zip": inner.getvalue()})

    parsed = [p for p in ing.parse_directory(str(tmp_path), include_pdf=False) if p.records]
    assert parsed and parsed[0].records[0]["Amount"] == "99"


def test_archive_skips_macos_noise(tmp_path):
    _zip_with(tmp_path, "a.zip", {
        "real.csv": "Amount,Date\n10,01-08-2024\n",
        "__MACOSX/._real.csv": "\x00\x05\x16\x07",
        "._real.csv": "\x00\x05\x16\x07",
    })
    parsed = ing.parse_directory(str(tmp_path), include_pdf=False)
    assert len(parsed) == 1


def test_archive_refuses_path_traversal(tmp_path):
    """A crafted member must not be written outside the extraction directory."""
    from backend.app.ingestion.parsers import archive

    _zip_with(tmp_path, "evil.zip", {"../escaped.csv": "Amount\n1\n"})
    dest = tmp_path / "out"
    extracted = archive.extract_archive(str(tmp_path / "evil.zip"), dest,
                                        max_total_bytes=1 << 20)

    assert extracted == []
    assert not (tmp_path / "escaped.csv").exists()


def test_archive_respects_expansion_budget(tmp_path):
    """A zip bomb must stop at the budget rather than filling the disk."""
    from backend.app.ingestion.parsers import archive

    _zip_with(tmp_path, "big.zip", {"a.csv": "x" * 5000, "b.csv": "y" * 5000})
    extracted = archive.extract_archive(str(tmp_path / "big.zip"), tmp_path / "o",
                                        max_total_bytes=6000)
    assert len(extracted) == 1  # first fits, second exceeds the remaining budget


def test_scratch_dirname_is_bounded_and_collision_resistant():
    """Two archives sharing a long prefix must not extract over each other."""
    from backend.app.ingestion.parsers import archive

    long_stem = "Fw__DO_DEBIT_FREEZE_and_also_provide_all_details_of_below_mentions" * 3
    a = archive.scratch_dirname(long_stem, "/case/a/x.zip")
    b = archive.scratch_dirname(long_stem, "/case/b/x.zip")

    assert len(a) <= archive._MAX_SCRATCH_COMPONENT + 9
    assert a != b, "same prefix, different archive — names must not collide"
    assert a == archive.scratch_dirname(long_stem, "/case/a/x.zip"), "must be stable"


def test_nested_archive_members_survive_long_source_names(tmp_path):
    """A long archive name must not push nested members past the 260-char Windows limit.

    Operators mail statements as a reply, so the mail subject becomes the filename: real
    `FIR 65-2024` archives run to 145 characters. The scratch directory was named after
    the archive at *every* nesting level, so outer + `__nested` + the member's own path
    crossed 260 and the members came back `[WinError 206] The filename or extension is
    too long`. Evidence lost to a path length we imposed on ourselves — and recorded as
    "unreadable", which reads like a corrupt exhibit.

    Asserted as a length invariant rather than via the OS, so it also fails on a platform
    that happens to allow the long path.
    """
    import io
    import shutil
    import tempfile
    import zipfile
    from pathlib import Path

    from backend.app.ingestion.parsers import archive

    # The real shape: a ~145-char reply-mail name. Both levels are given one, so that a
    # regression at *either* the outer (service._walk) or the nested (extract_archive)
    # naming blows the budget on its own — with only one of them long, the other's
    # bounding hides the fault and this test passes while the bug is live.
    outer_stem = ("Fw__DO_DEBIT_FREEZE_and_also_provide_all_details_of_below_mentions_"
                  "Debit_card_which_is__linked_with_bank_accounts_Regarding_CCPS_FIR_0065")
    # Identifiers are synthetic placeholders of the same *length* as the real ones — the
    # length is the whole point of this fixture, and rule 4 keeps case numbers out of git.
    inner_stem = ("Re__ACCOUNTNOREDACT_KYC_PERMANENT_ADDRESS_PROOF_and_TERMS_CONDITION_"
                  "FORM_and_PANCARD_FORM_60_61_49A_regarding_CCPS_FIR_0065_2024_reply")

    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as inner:
        inner.writestr("ACCOUNTNOREDACT KYC/E_SIGNATURE.csv", "Amount,Date\n10,01-08-2024\n")

    outer = tmp_path / f"{outer_stem}.zip"
    with zipfile.ZipFile(outer, "w") as zf:
        zf.writestr(f"{inner_stem}.zip", buf.getvalue())

    # Extract under a real mkdtemp root, exactly as service._walk does. pytest's own
    # tmp_path is ~90 characters, which would spend the budget on the harness rather
    # than on the geometry under test.
    scratch = Path(tempfile.mkdtemp(prefix="erakshak-archives-"))
    try:
        dest = scratch / archive.scratch_dirname(outer.stem, str(outer))
        extracted = archive.extract_archive(str(outer), dest, max_total_bytes=1 << 20)

        assert any(f.name == "E_SIGNATURE.csv" for f in extracted), (
            "nested member was lost; before the fix this raised WinError 206")

        # The limit Windows actually enforces, on the absolute path.
        for f in extracted:
            assert len(str(f)) < 260, f"path too long ({len(str(f))} chars): {f}"
    finally:
        shutil.rmtree(scratch, ignore_errors=True)


def test_member_names_windows_cannot_represent_are_still_recovered(tmp_path):
    """A member named with a colon must not be lost on Windows.

    ZIPs written on Linux/macOS carry characters Windows refuses. Real evidence does: a
    Bandhan reply in `FIR 65-2024` ships members under
    `summary_BNB-REFNUM_FIR 0065:2024_17_06_2025`, and every one failed with
    `[WinError 267] The directory name is invalid` — logged as "member unreadable", which
    reads like a corrupt exhibit rather than a limit of our own filesystem.
    """
    from backend.app.ingestion.parsers import archive

    _zip_with(tmp_path, "reply.zip", {
        "summary_BNB-REFNUM_FIR 0065:2024_17_06_2025/report.csv": "Amount,Date\n10,01-08-2024\n",
    })
    extracted = archive.extract_archive(str(tmp_path / "reply.zip"), tmp_path / "o",
                                        max_total_bytes=1 << 20)

    assert [f.name for f in extracted] == ["report.csv"]
    assert extracted[0].read_text(encoding="utf-8").startswith("Amount,Date")


def test_sanitising_member_names_does_not_weaken_the_traversal_guard():
    """`..` must survive sanitisation, or the path-escape check stops seeing an escape."""
    from backend.app.ingestion.parsers import archive

    assert archive._sanitise_for_windows("../escaped.csv") == "../escaped.csv"
    assert archive._sanitise_for_windows("a/../../b.csv") == "a/../../b.csv"
    assert archive._sanitise_for_windows("FIR 0065:2024/x.csv") == "FIR 0065_2024/x.csv"


def test_encrypted_archive_does_not_abort_the_batch(tmp_path):
    """Operators send password-protected archives; zipfile raises a bare RuntimeError.

    One locked member must not lose the rest of the archive — and must never abort the
    whole case, which is what happened before this was handled per-member.
    """
    import zipfile

    from backend.app.ingestion.parsers import archive

    p = tmp_path / "locked.zip"
    with zipfile.ZipFile(p, "w") as zf:
        zf.writestr("open.csv", "Amount\n1\n")
    # Mark one member encrypted without a real password (flag bit 0x1).
    with zipfile.ZipFile(p, "a") as zf:
        zf.writestr("secret.csv", "Amount\n2\n")
    with zipfile.ZipFile(p, "a") as zf:
        zf.infolist()[-1].flag_bits |= 0x1

    extracted = archive.extract_archive(str(p), tmp_path / "o", max_total_bytes=1 << 20)
    assert any(f.name == "open.csv" for f in extracted)


# ── Word documents hold many tables, not one ───────────────────────────────────

def _docx_with_tables(path, tables):
    from docx import Document
    doc = Document()
    for rows in tables:
        t = doc.add_table(rows=len(rows), cols=len(rows[0]))
        for r, row in enumerate(rows):
            for c, val in enumerate(row):
                t.cell(r, c).text = val
    doc.save(path)


def test_read_all_grids_returns_every_table(tmp_path):
    from backend.app.ingestion.parsers import docx_tables

    p = tmp_path / "ca_report.docx"
    _docx_with_tables(p, [
        [["Account Number", "IFSC"], ["111", "AAA"]],
        [["Account Number", "IFSC"], ["222", "BBB"]],
        [["Account Number", "IFSC"], ["333", "CCC"]],
    ])
    grids = docx_tables.read_all_grids(str(p))
    assert len(grids) == 3


def test_multi_table_docx_yields_one_parsedfile_per_table(tmp_path):
    """Keeping only the largest table discarded 2,540 of 5,430 real rows."""
    p = tmp_path / "ca_report.docx"
    _docx_with_tables(p, [
        [["Account Number", "IFSC"], ["111", "AAA"]],
        [["Account Number", "IFSC"], ["222", "BBB"]],
    ])
    parsed = ing.parse_file_multi(str(p))

    assert len(parsed) == 2
    assert [pf.table_index for pf in parsed] == [1, 2]
    values = sorted(r["Account Number"] for pf in parsed for r in pf.records)
    assert values == ["111", "222"]


# ── Finacle/IndusInd bulk SOA: FORACID + DEDIT_AMOUNT ──────────────────────────
# statement bulk.xls (6,975 rows on fir-65-2024) matched bank_generic then lost
# every row: the account lived in FORACID (not an alias) and debit was misspelled
# DEDIT_AMOUNT. Timestamp parsed fine — reject reason was "no account".

FINACLE_BULK_CSV = (
    "FORACID,ACCT_NAME,TRAN_DATE,TRAN_PARTICULAR,DEDIT_AMOUNT,CREDIT_AMOUNT,TRAN_ID\n"
    "100240778506,BAMBHANIYA A,09-05-24,UPI/CR/okaxis,0,20249,S90294351\n"
    "100240778506,BAMBHANIYA A,10-05-24,UPI/DR/okaxis,500,0,S90294352\n"
)


def test_finacle_foracid_bulk_statement_survives_normalization(tmp_path):
    from backend.app.normalization import service as norm

    p = tmp_path / "statement_bulk.csv"
    p.write_text(FINACLE_BULK_CSV, encoding="utf-8")
    pf = ing.parse_file(str(p))
    assert pf.source_type == "BANK"
    assert len(pf.records) == 2

    events, rejects = norm.normalize_parsed_files([pf])
    assert len(events) == 2, rejects
    assert events[0]["primary"] == ("ACCOUNT_NO", "100240778506")
    assert events[0]["direction"] == "CREDIT" and events[0]["amount"] == 20249.0
    assert events[1]["direction"] == "DEBIT" and events[1]["amount"] == 500.0


# ── Exchange wallet ledger: Time + User ID + signed Amount ─────────────────────
# wallet_details / BNB reports (6,683 rejected rows) matched bank_generic at 0.25
# via Description, then every row lost its timestamp because `Time` was not a
# date alias and `User ID` was not an account alias.

EXCHANGE_LEDGER_CSV = (
    "Transaction ID,User ID,Currency,Type,Amount,Available,Description,Time\n"
    "256506709293,1020260104,USDT,Product withdrawal success,-19,0,"
    "Automatically confirm,2025-04-21 14:04:10\n"
    "243255041159,1014659199,USDC,Future transfer,0.1082099,0.1082099,"
    "Clear Zombie users,2025-02-16 21:50:49\n"
)


def test_exchange_wallet_ledger_survives_normalization(tmp_path):
    from backend.app.normalization import service as norm

    p = tmp_path / "wallet_details.csv"
    p.write_text(EXCHANGE_LEDGER_CSV, encoding="utf-8")
    pf = ing.parse_file(str(p))
    assert pf.source_type == "BANK"
    assert pf.profile_id in {"crypto_exchange_ledger", "bank_generic"}

    events, rejects = norm.normalize_parsed_files([pf])
    assert len(events) == 2, (rejects, pf.profile_id, pf.headers)
    assert events[0]["primary"][0] == "ACCOUNT_NO"
    assert events[0]["direction"] == "DEBIT" and events[0]["amount"] == 19.0
    assert events[0]["asset"] == "CRYPTO:USDT"
    assert events[1]["direction"] == "CREDIT"
    assert events[1]["asset"] == "CRYPTO:USDC"


# ── Tab-separated IPDR range export ────────────────────────────────────────────
# ipdr__1365.txt is the same schema as the working .xlsx, but tab-delimited.
# Default comma sep kept the header as one column → unrecognized source, 9/9 lost.

IPDR_TAB_TXT = (
    "IP\tVALUE\tF DATE\tF TIME\tT DATE\tT TIME\n"
    "IPV6\t2409:40d2:1328:a31c::1\t20241226\t135728\t20241226\t140128\n"
    "IPV6\t2409:40d2:132d:b3f8::2\t20241125\t132132\t20241125\t132532\n"
)


def test_tab_separated_ipdr_txt_is_recognized(tmp_path):
    from backend.app.normalization import service as norm

    p = tmp_path / "ipdr_range.txt"
    p.write_text(IPDR_TAB_TXT, encoding="utf-8")
    pf = ing.parse_file(str(p))
    assert pf.source_type == "IPDR", (pf.source_type, pf.headers, pf.profile_id)
    assert len(pf.records) == 2

    events, rejects = norm.normalize_parsed_files([pf])
    assert len(events) == 2, rejects
    assert events[0]["event_type"] == "IP_SESSION"


def test_registered_mobile_survives_country_code_separator():
    """"Mobile: +91 8180934367" must yield the whole number, not "+91".

    The old digits-only pattern stopped at the first space, so three of four real
    statements extracted the bare country code. A country code identifies nobody, and
    each one lost is a lost account-to-phone bridge — the exact link FR-9 needs.
    """
    from backend.app.ingestion.service import _extract_identity

    for text, expect in [
        ("Registered Mobile: +91 8180934367", "+918180934367"),
        ("Registered Mobile : +91-81809 34367", "+918180934367"),
        ("Mobile Number: 8180934367", "8180934367"),
        ("Customer Mobile: 0 8180934367", "08180934367"),
    ]:
        got = _extract_identity([], [text])
        assert got.get("registered_mobile") == expect, f"{text!r} -> {got}"

    # a country code with no number must be refused, not stored
    assert "registered_mobile" not in _extract_identity([], ["Registered Mobile: +91"])
    assert "registered_mobile" not in _extract_identity([], ["Registered Mobile: -"])

    # split across grid cells: ["Registered Mobile", "+91", "8180934367"]
    grid = [["Registered Mobile", "+91", "8180934367"]]
    assert _extract_identity(grid, []).get("registered_mobile") == "+918180934367"

    # names must be unaffected by the numeric path
    assert _extract_identity([], ["Account Name: Smita Gawali"]).get("account_holder")


def test_stacked_tables_in_one_grid_are_split():
    """Two tables pasted into one sheet must become two sections, not one.

    Portal and complaint exports stack unrelated tables in a single grid. With one
    header row for the whole grid, every row below the second table's header is mapped
    against the first table's columns.

    Deliberately conservative: a false split severs a real table, which is worse than
    leaving a section unrecovered — so a candidate header needs several known-alias
    hits and must be mostly non-numeric.
    """
    from backend.app.ingestion.service import _split_grid_sections

    grid = [
        ["Tran Date", "Narration", "Debit", "Credit", "Balance"],
        ["01/04/2024", "UPI/ABC", "", "5000", "5000"],
        ["02/04/2024", "UPI/DEF", "1000", "", "4000"],
        ["Transaction Date", "Account No./ Wallet ID", "Transaction Amount"],
        ["03/04/2024", "259024319039", "2,00,000.00"],
        ["04/04/2024", "201029737717", "5,50,000.00"],
    ]
    sections = _split_grid_sections(grid)
    assert len(sections) == 2, sections
    assert sections[0] == (0, 3)
    assert sections[1] == (3, 6)

    # a single-table grid must stay single — this is the no-behaviour-change guarantee
    assert len(_split_grid_sections(grid[:3])) == 1

    # numeric rows must never be mistaken for headers, however many columns they fill
    numeric_only = [
        ["Tran Date", "Narration", "Debit", "Credit", "Balance"],
        ["01/04/2024", "UPI/ABC", "1", "2", "3"],
        ["02/04/2024", "UPI/DEF", "4", "5", "6"],
    ]
    assert len(_split_grid_sections(numeric_only)) == 1

    assert _split_grid_sections([]) == []


def test_profile_may_claim_a_file_it_can_demonstrably_map():
    """A parseable statement must not be refused because required_any drifted.

    `match.required_any` and `field_map.aliases` are separate lists that must agree.
    A real statement headed "Trans Date and Time | Transaction Details | Debit |
    Credit | Balance" maps six canonical targets cleanly, yet scored 0.0 because
    required_any wanted the exact string "debit amount".
    """
    from backend.app.ingestion import detector

    statement = ["Trans Date and\nTime", "Value Date", "Transaction Details",
                 "Cheque No", "Debit", "Credit", "Balance"]
    got = detector.detect_profile(statement)
    assert got["source"] == "BANK", got
    assert got["confidence"] > 0.0
    # inferred, not asserted — the analyst must see it needs review
    assert got["confidence"] <= 0.49
    assert got["needs_manual_mapping"] is True

    # A genuine required_any match must still outrank a fallback and keep full confidence.
    proper = detector.detect_profile(
        ["Tran Date", "Narration", "Debit", "Credit", "Balance", "Ac_No"])
    assert proper["confidence"] > got["confidence"]

    # Too little to be sure: a lone date column claims nothing.
    assert detector.detect_profile(["Value Date"])["confidence"] == 0.0
    # No time anchor -> no fallback, however many other columns match.
    assert detector.detect_profile(["Debit", "Credit", "Balance"])["confidence"] == 0.0
    assert detector.detect_profile([])["confidence"] == 0.0


def test_required_all_stays_a_hard_gate():
    """required_all means "this shape is mandatory" — no fallback may bypass it."""
    from backend.app.ingestion import detector

    profile = {
        "profile": {"id": "t", "source": "BANK",
                    "match": {"required_all": ["mandatory col"]}},
        "field_map": {
            "timestamp_start": {"aliases": ["Tran Date"]},
            "account_no": {"aliases": ["Ac_No"]},
            "credit": {"aliases": ["Credit"]},
        },
    }
    # maps three targets incl. time + subject, but the mandatory column is absent
    assert detector.score_profile(["Tran Date", "Ac_No", "Credit"], profile) == 0.0
    assert detector.score_profile(
        ["Tran Date", "Ac_No", "Credit", "Mandatory Col"], profile) > 0.0


def test_a_duplicated_exhibit_is_parsed_once_and_recorded(tmp_path):
    """One portal export appeared three times across `fir-0006-2025-u` and its 830 rows
    were parsed three times. Event-level dedup meant the output was never wrong, but the
    work was wasted and the reject counts read far worse than the evidence.

    The copy is recorded rather than ignored: which exhibits are duplicated is part of the
    chain of custody.
    """
    from backend.app.ingestion import service as ingestion

    body = ("Tran Date,Tran Particular,Debit Amount,Credit Amount\n"
            "01-02-2024,UPI/1234/PAY,,5000.00\n"
            "02-02-2024,ATM WDL SURAT,2000.00,\n"
            "03-02-2024,NEFT-VENDOR,1000.00,\n"
            "04-02-2024,IMPS/9999/REFUND,,500.00\n")
    (tmp_path / "a").mkdir()
    (tmp_path / "b").mkdir()
    (tmp_path / "a" / "stmt.csv").write_text(body, encoding="utf-8")
    (tmp_path / "b" / "stmt_copy.csv").write_text(body, encoding="utf-8")
    (tmp_path / "b" / "other.csv").write_text(
        body.replace("5000.00", "6000.00"), encoding="utf-8")

    skipped: list[dict] = []
    parsed = ingestion.parse_directory(str(tmp_path), skipped_out=skipped)

    assert len(parsed) == 2, "byte-identical copy should not be parsed twice"
    dupes = [s for s in skipped if s.get("duplicate_of")]
    assert len(dupes) == 1
    assert dupes[0]["duplicate_of"] == "stmt.csv"
    assert "byte-identical" in dupes[0]["reason"]
